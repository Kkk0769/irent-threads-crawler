"""Simple Word report: numbered comment, then its clickable URL."""
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


def clean_text(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)


def export_word(rows, path, status, scanned_count):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ['Normal', 'Title']:
        style = doc.styles[name]
        style.font.name = 'Microsoft JhengHei'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.25
    doc.styles['Title'].font.size = Pt(20)
    doc.add_paragraph('iRent 負面評論整理', style='Title')
    doc.add_paragraph(f'讀取 {scanned_count} 篇貼文，篩出 {len(rows)} 篇疑似負評。{status}。')
    doc.add_paragraph('依找到的順序排列。評論為頁面取得的文字，可能未展開全文；請由連結確認內容與語意。')
    for index, row in enumerate(rows, 1):
        content = clean_text(row.get('title_text') or row['text'])
        paragraph = doc.add_paragraph(f'{index}. {content}')
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_after = Pt(3)
        link_paragraph = doc.add_paragraph()
        link_paragraph.paragraph_format.space_after = Pt(16)
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), doc.part.relate_to(row['permalink'], RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '0563C1')
        props.append(color)
        run.append(props)
        text = OxmlElement('w:t')
        text.text = row['permalink']
        run.append(text)
        hyperlink.append(run)
        link_paragraph._p.append(hyperlink)
    if not rows:
        doc.add_paragraph('本次未篩出疑似負評，不代表 Threads 上沒有相關評論。')
    doc.save(path)
